"""Endpoints de reportes exportables (Excel/PDF)."""

from io import BytesIO
from datetime import datetime

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from sqlalchemy import func
from sqlalchemy.orm import Session

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from app.api.deps import get_current_tenant
from app.core.database import get_db
from app.models import StockLot, Tenant, Warehouse, Resource

router = APIRouter(prefix="/reports", tags=["reports"])


def _get_stock_data(db: Session):
    """Obtiene los datos de stock consolidado."""
    query = (
        db.query(
            Warehouse.name.label("warehouse_name"),
            Resource.name.label("resource_name"),
            func.sum(StockLot.quantity).label("total_quantity"),
            func.count(StockLot.id).label("lot_count"),
            func.min(StockLot.expiration).label("nearest_expiration"),
        )
        .join(Warehouse, StockLot.warehouse_id == Warehouse.id)
        .join(Resource, StockLot.resource_id == Resource.id)
        .group_by(Warehouse.name, Resource.name)
    )
    return query.all()


@router.get("/stock/excel")
def export_stock_excel(
    current_tenant: Tenant = Depends(get_current_tenant),
    db: Session = Depends(get_db),
):
    """Exporta el stock consolidado a un archivo Excel."""
    rows = _get_stock_data(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Consolidado"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF")
    
    from openpyxl.styles import PatternFill
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    # Cabeceras
    headers = ["Almacén", "Recurso", "Cantidad Total", "Lotes", "Próxima Caducidad"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Filas
    for row_num, r in enumerate(rows, 2):
        ws.cell(row=row_num, column=1, value=r.warehouse_name)
        ws.cell(row=row_num, column=2, value=r.resource_name)
        ws.cell(row=row_num, column=3, value=float(r.total_quantity))
        ws.cell(row=row_num, column=4, value=r.lot_count)
        ws.cell(row=row_num, column=5, value=r.nearest_expiration.strftime("%Y-%m-%d") if r.nearest_expiration else "—")

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = max_length + 2

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=stock_{current_tenant.schema_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"},
    )


@router.get("/stock/pdf")
def export_stock_pdf(
    current_tenant: Tenant = Depends(get_current_tenant),
    db: Session = Depends(get_db),
):
    """Exporta el stock consolidado a un archivo PDF."""
    rows = _get_stock_data(db)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elements = []

    # Estilos
    styles = getSampleStyleSheet()
    title_style = styles['Heading1']
    normal_style = styles['Normal']

    # Título
    elements.append(Paragraph(f"Reporte de Stock - {current_tenant.name}", title_style))
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}", normal_style))
    elements.append(Spacer(1, 12))

    # Tabla
    data = [["Almacén", "Recurso", "Cantidad", "Lotes", "Próx. Caducidad"]]
    for r in rows:
        data.append([
            r.warehouse_name,
            r.resource_name,
            f"{float(r.total_quantity):,.2f}",
            str(r.lot_count),
            r.nearest_expiration.strftime("%Y-%m-%d") if r.nearest_expiration else "—",
        ])

    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F81BD")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=stock_{current_tenant.schema_name}_{datetime.now().strftime('%Y%m%d')}.pdf"},
    )